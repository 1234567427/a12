import streamlit as st
from pptx import Presentation
import pdfplumber
import io
import json
import asyncio
import re
import base64
import edge_tts
from openai import OpenAI
import os
os.environ.pop("HTTP_PROXY", None)
os.environ.pop("HTTPS_PROXY", None)
os.environ.pop("ALL_PROXY", None)
os.environ.pop("http_proxy", None)
os.environ.pop("https_proxy", None)
os.environ.pop("all_proxy", None)

API_KEY = st.secrets["ZHIPU_API_KEY"]
BASE_URL = "https://open.bigmodel.cn/api/paas/v4/"
MODEL = "glm-4-flash"

client = OpenAI(api_key=API_KEY, base_url=BASE_URL)

st.title("A12 智课生成系统")

uploaded = st.file_uploader("上传课件", type=["pptx", "pdf"])

def split_sentences(text):
    parts = re.split(r'(?<=[。！？；])', text)
    return [p.strip() for p in parts if p.strip()]

def extract_pptx(file):
    prs = Presentation(file)
    text = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t:
                        text.append(t)
    return "\n".join(text)

def extract_pdf(file):
    text = []
    with pdfplumber.open(file) as pdf:
        for page in pdf.pages:
            t = page.extract_text()
            if t:
                text.append(t)
    return "\n".join(text)

def get_knowledge_points(raw_text):
    prompt = "你是一个教学助手。请阅读下面的课件内容，提取出知识点列表。\n要求：\n1. 每个知识点包含：标题、简要说明、重要程度（高/中/低）\n2. 按教学逻辑顺序排列\n3. 只输出 JSON 数组，不要输出其他内容\n\n格式示例：\n[\n  {\"title\": \"函数定义\", \"desc\": \"介绍函数的基本概念\", \"level\": \"高\"}\n]\n\n课件内容：\n" + raw_text[:3000]
    response = client.chat.completions.create(
        model=MODEL,
        messages=[{"role": "user", "content": prompt}],
        temperature=0.3
    )
    return response.choices[0].message.content

def get_script(knowledge_points):
    prompt = "你是一位老师，请根据下面的知识点列表，生成一段讲课脚本。\n要求：\n1. 包含开场白、每个知识点的讲解、知识点之间的过渡语、结束语\n2. 语言口语化，像在课堂上讲课\n3. 每个知识点的讲解控制在 100 字左右\n4. 只输出 JSON 数组，每个元素包含 title 和 script 两个字段\n\n格式示例：\n[\n  {\"title\": \"开场白\", \"script\": \"同学们好，今天我们来学习...\"}\n]\n\n知识点列表：\n" + knowledge_points
    response = client.chat.completions.create(
        model=MODEL,
        messages=[{"role": "user", "content": prompt}],
        temperature=0.5
    )
    return response.choices[0].message.content

def answer_question(question, context, history):
    messages = [
        {"role": "system", "content": "你是一位课程助教。请根据下面的课件内容回答学生问题。如果课件里没有相关内容，就说“这个问题课件里没有提到”。回答要简洁、口语化。\n\n课件内容：\n" + context[:3000]}
    ]
    for h in history:
        messages.append(h)
    messages.append({"role": "user", "content": question})
    response = client.chat.completions.create(
        model=MODEL,
        messages=messages,
        temperature=0.5
    )
    return response.choices[0].message.content

def judge_topic(question, knowledge_points):
    prompt = "下面是一份知识点列表和一个学生问题。请判断学生的问题最可能对应哪个知识点，只输出该知识点的标题，不要输出其他内容。如果无法判断，输出“无”。\n\n知识点列表：\n" + knowledge_points + "\n\n学生问题：" + question
    response = client.chat.completions.create(
        model=MODEL,
        messages=[{"role": "user", "content": prompt}],
        temperature=0.1
    )
    return response.choices[0].message.content.strip()

async def text_to_speech(text, filename):
    communicate = edge_tts.Communicate(text, "zh-CN-XiaoxiaoNeural")
    await communicate.save(filename)

def generate_audio(text, filename):
    asyncio.run(text_to_speech(text, filename))

def build_player(audio_files, sentences):
    audio_data = []
    for f in audio_files:
        try:
            with open(f, "rb") as fp:
                b64 = base64.b64encode(fp.read()).decode()
                audio_data.append("data:audio/mp3;base64," + b64)
        except:
            audio_data.append("")

    audio_json = json.dumps(audio_data)
    sentence_json = json.dumps(sentences)

    html = """
    <html>
    <body>
    <button id="playBtn" style="font-size:16px;padding:8px 16px;margin-bottom:10px;">▶ 开始播放</button>
    <div id="subs" style="font-size:16px;line-height:2;"></div>
    <script>
    var audios = AUDIO_JSON;
    var sentences = SENTENCE_JSON;
    var current = 0;
    var audio = new Audio();
    audio.src = audios[0];

    var subsDiv = document.getElementById("subs");
    var playBtn = document.getElementById("playBtn");

    sentences.forEach(function(s, i){
        var span = document.createElement("span");
        span.innerText = s;
        span.style.cursor = "pointer";
        span.style.marginRight = "6px";
        span.style.padding = "2px";
        span.onclick = function(){
            current = i;
            audio.src = audios[i];
            audio.play();
            playBtn.innerText = "⏸ 暂停";
            highlight();
        };
        subsDiv.appendChild(span);
    });

    function highlight(){
        var spans = subsDiv.getElementsByTagName("span");
        for(var i=0;i<spans.length;i++){
            spans[i].style.background = (i===current) ? "#ffe58f" : "transparent";
        }
    }

    audio.onended = function(){
        if(current < audios.length - 1){
            current += 1;
            audio.src = audios[current];
            audio.play();
            highlight();
        } else {
            playBtn.innerText = "▶ 开始播放";
        }
    };

    playBtn.onclick = function(){
        if(audio.paused){
            audio.play();
            playBtn.innerText = "⏸ 暂停";
        } else {
            audio.pause();
            playBtn.innerText = "▶ 继续播放";
        }
    };

    highlight();
    </script>
    </body>
    </html>
    """
    html = html.replace("AUDIO_JSON", audio_json).replace("SENTENCE_JSON", sentence_json)
    return html

if uploaded:
    st.success("已上传：" + uploaded.name)
    file_bytes = uploaded.read()
    file_obj = io.BytesIO(file_bytes)

    if uploaded.name.endswith(".pptx"):
        raw_text = extract_pptx(file_obj)
    else:
        raw_text = extract_pdf(file_obj)

    st.session_state["raw_text"] = raw_text

    st.subheader("原始文本")
    st.text_area("内容", raw_text, height=150)

    if st.button("提取知识点"):
        with st.spinner("AI 正在分析..."):
            result = get_knowledge_points(raw_text)
        st.session_state["knowledge"] = result
        st.subheader("知识点列表")
        st.write(result)

    if "knowledge" in st.session_state:
        if st.button("生成讲授脚本"):
            with st.spinner("AI 正在写讲稿..."):
                script = get_script(st.session_state["knowledge"])
            st.session_state["script"] = script
            st.subheader("讲授脚本")
            st.write(script)

    if "script" in st.session_state:
        if st.button("生成语音"):
            with st.spinner("正在合成语音..."):
                try:
                    script_list = json.loads(st.session_state["script"])
                except:
                    st.error("脚本解析失败，请重新生成")
                    st.stop()
                for i, item in enumerate(script_list):
                    sentences = split_sentences(item["script"])
                    for j, sent in enumerate(sentences):
                        filename = "audio_" + str(i) + "_" + str(j) + ".mp3"
                        generate_audio(sent, filename)

        try:
            script_list = json.loads(st.session_state["script"])
        except:
            st.stop()

        st.subheader("讲授内容（点击句子跳转，自动连播）")

        all_audios = []
        all_sentences = []
        for i, item in enumerate(script_list):
            sentences = split_sentences(item["script"])
            for j, sent in enumerate(sentences):
                all_audios.append("audio_" + str(i) + "_" + str(j) + ".mp3")
                all_sentences.append(sent)

        html = build_player(all_audios, all_sentences)
        st.components.v1.html(html, height=400)

st.divider()
st.subheader("课堂问答")

if "chat_history" not in st.session_state:
    st.session_state["chat_history"] = []

question = st.text_input("输入你的问题")

if st.button("提问"):
    if question:
        with st.spinner("AI 正在回答..."):
            context = st.session_state.get("raw_text", "")
            answer = answer_question(question, context, st.session_state["chat_history"])
            st.session_state["chat_history"].append({"role": "user", "content": question})
            st.session_state["chat_history"].append({"role": "assistant", "content": answer})
            if "knowledge" in st.session_state:
                topic = judge_topic(question, st.session_state["knowledge"])
                st.session_state["last_topic"] = topic

for msg in st.session_state["chat_history"]:
    if msg["role"] == "user":
        st.markdown("**你：** " + msg["content"])
    else:
        st.markdown("**AI：** " + msg["content"])

if "last_topic" in st.session_state:
    st.info("你刚才问的是「" + st.session_state["last_topic"] + "」")

    if st.button("从这里继续讲授"):
        topic = st.session_state["last_topic"]
        try:
            script_list = json.loads(st.session_state["script"])
        except:
            st.error("请先生成讲授脚本")
            st.stop()

        start_idx = 0
        for i, item in enumerate(script_list):
            if topic in item["title"] or item["title"] in topic:
                start_idx = i
                break

        st.success("从「" + script_list[start_idx]["title"] + "」继续讲授")

        cont_audios = []
        cont_sentences = []
        for i in range(start_idx, len(script_list)):
            item = script_list[i]
            sentences = split_sentences(item["script"])
            for j, sent in enumerate(sentences):
                cont_audios.append("audio_" + str(i) + "_" + str(j) + ".mp3")
                cont_sentences.append(sent)

        html = build_player(cont_audios, cont_sentences)
        st.components.v1.html(html, height=400)